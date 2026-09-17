"""SETUP STAGE — RESUME -> FACT BANK. An existing .docx/.pdf resume -> markdown.

The fact bank is the single source of truth the tailor stage may draw on and the
grounding check verifies against. Building it by hand from an existing resume is
the most tedious part of setting this app up, so this stage does the first pass.

The safety argument is the whole design here. Everything downstream trusts this
file completely: if the extractor embellishes, the grounding check will happily
certify the embellishment as "supported" and the fabrication reaches a real
resume sent to a real employer. So:

  1. EXTRACTION IS NOT GENERATION. The prompt forbids adding, inferring, or
     upgrading verbs, and the temperature is 0.
  2. UNCERTAINTY IS MARKED, not smoothed over. Anything vague in the source comes
     back flagged so you can see what to verify.
  3. A HUMAN ALWAYS REVIEWS. This stage returns markdown; it never writes
     data/fact_bank.md. The Setup pane shows it in an editor and you press Save.

Same LLM/Python split as the tailor stage: the model returns STRUCTURE
(FactBankDoc), Python renders the markdown. The model never emits the artefact.

LEARN: extraction vs generation, why a trusted-source file needs a stricter
       prompt than a draft, human-in-the-loop on ingestion (not just output).
"""
from __future__ import annotations

import io
import json
import re

from schemas import FactBankDoc
from llm import call_structured

_SCHEMA = json.dumps(FactBankDoc.model_json_schema(), indent=2)

SYSTEM = f"""You convert a candidate's EXISTING resume into a factual "fact bank":
a plain inventory of what this person has actually done.

This file becomes the ONLY source a resume writer may draw on, and a separate
fact-checker will verify every future claim against it. Anything you add here that
the resume did not say becomes a lie on a real job application.

Rules:
- Extract ONLY what the resume states. Never infer, never generalise, never add a
  tool, metric, date, or responsibility that is not written down.
- Preserve the exact strength of every verb. If it says "used", write "used" — not
  "built", "designed", or "led". Do not upgrade "familiar with" into "proficient".
- Keep concrete specifics verbatim: numbers, dataset names, library names, dates,
  team sizes, percentages. These are what make a grounded resume credible.
- Do NOT copy marketing prose ("passionate self-starter"). Facts only.
- Set `uncertain: true` on any item where the resume was vague, ambiguous, or you
  had to guess at the meaning. Flagging is always better than guessing silently.
- Preserve the resume's own grouping as `sections` (Experience, Projects, Skills,
  Education, and so on). Use the headings the resume actually uses.
- If the resume mentions gaps or tools the person is still learning, keep them in a
  section of their own. Honest gaps are load-bearing: the writer is instructed never
  to claim proficiency in them.

Return ONLY a JSON object conforming to this schema (no prose, no markdown fences):
{_SCHEMA}
"""

# The Setup pane explains the review step; this header survives into the saved
# file so the reminder is still there weeks later when the file is edited by hand.
_HEADER = """# Fact Bank

Extracted from an uploaded resume. REVIEW THIS BEFORE RELYING ON IT — the resume
tailor may use only what is written here, and the grounding check treats every
line as true. Anything marked `(verify)` was vague in the source document.
"""


def extract_docx_text(data: bytes) -> str:
    """Paragraphs + table cells from a .docx. Tables matter: plenty of resume
    templates lay the whole document out in an invisible table, and reading only
    paragraphs returns an almost empty string for those."""
    import docx                                  # already a dependency (render.py)

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(p.strip() for p in parts if p and p.strip())


def extract_pdf_text(data: bytes) -> str:
    """Text from a PDF. Raises a readable error when the file is a scan — pypdf
    returns "" for image-only pages, and an empty fact bank that looks like a
    successful extraction is exactly the silent failure this project avoids."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    if not text.strip():
        raise ValueError(
            "No text found in that PDF. It is probably a scan or an image export "
            "— save it as a text PDF or upload the .docx instead."
        )
    return text


def extract_text(data: bytes, ext: str) -> str:
    """Dispatch on extension. Mirrors the .docx/.pdf gate the resume-upload
    endpoint already enforces (api.py upload_resume)."""
    ext = ext.lower()
    if ext == ".docx":
        return extract_docx_text(data)
    if ext == ".pdf":
        return extract_pdf_text(data)
    raise ValueError(f"Unsupported file type {ext!r}. Upload a .docx or .pdf.")


def render_markdown(doc: FactBankDoc) -> str:
    """FactBankDoc -> the markdown actually saved. Deterministic Python, so the
    same extraction always produces the same file and a diff is meaningful."""
    out = [_HEADER]
    if doc.name:
        out.append(f"**{doc.name}**")
    if doc.contact:
        out.append(doc.contact)
    if doc.name or doc.contact:
        out.append("")

    for section in doc.sections:
        out.append(f"## {section.heading}")
        for item in section.items:
            suffix = "  *(verify)*" if item.uncertain else ""
            out.append(f"- {item.text}{suffix}")
        out.append("")
    return "\n".join(out).rstrip() + "\n"


# A resume is a couple of pages; this cap is a guard against someone uploading a
# 200-page PDF and sending the whole thing to the model at the quality tier.
MAX_CHARS = 60_000


def extract_fact_bank(text: str) -> str:
    """Resume text -> fact bank markdown, ready for human review.

    Quality tier at temperature 0: this is the strictest extraction in the app,
    and the cheap tier is markedly worse at resisting the urge to smooth vague
    phrasing into confident claims.
    """
    text = (text or "").strip()
    if not text:
        raise ValueError("That file contained no readable text.")
    text = _collapse_blank_lines(text)[:MAX_CHARS]

    doc: FactBankDoc = call_structured(
        SYSTEM, f"RESUME:\n{text}", schema=FactBankDoc,
        tier="quality", temperature=0.0)
    return render_markdown(doc)


def _collapse_blank_lines(text: str) -> str:
    """PDF and docx extraction both emit runs of blank lines; they cost tokens
    and tell the model nothing."""
    return re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n"))


def extract_from_upload(data: bytes, ext: str) -> str:
    """The endpoint's entry point: bytes -> reviewed-pending markdown."""
    return extract_fact_bank(extract_text(data, ext))
