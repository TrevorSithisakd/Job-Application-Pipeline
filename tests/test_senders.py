"""Tests for inbox sender discovery and the fact-bank extractor.

The Gmail client is faked at the service-object boundary — the same seam
`_gmail_service()` already provides — so pagination, header parsing, and the
per-message metadata fetch are all exercised for real against canned responses.
No network, no credentials, no mailbox.
"""
from __future__ import annotations

import pytest

import settings
from schemas import (EmailSource, FactBankDoc, FactItem, FactSection,
                     SenderClassification, SenderVerdict)
from stages import factbank, ingest


# --- fake Gmail --------------------------------------------------------------

class _Exec:
    def __init__(self, value): self._value = value
    def execute(self): return self._value


class _FakeMessages:
    """Two pages of results, so the pagination loop is actually covered."""

    def __init__(self, msgs, pages):
        self.msgs, self.pages = msgs, pages
        self.list_calls, self.get_calls = [], []

    def list(self, userId, q, maxResults=100, pageToken=None):
        self.list_calls.append({"q": q, "pageToken": pageToken})
        ids, nxt = self.pages[pageToken]
        body = {"messages": [{"id": i} for i in ids]}
        if nxt:
            body["nextPageToken"] = nxt
        return _Exec(body)

    def get(self, userId, id, format=None, metadataHeaders=None):
        self.get_calls.append({"id": id, "format": format,
                               "metadataHeaders": metadataHeaders})
        frm, subj = self.msgs[id]
        return _Exec({"payload": {"headers": [{"name": "From", "value": frm},
                                              {"name": "Subject", "value": subj}]}})


def _fake_service(msgs, pages):
    box = _FakeMessages(msgs, pages)
    service = type("S", (), {"users": lambda self: type("U", (), {
        "messages": lambda self: box})()})()
    return service, box


MSGS = {
    "1": ("SEEK <noreply@s.seek.com.au>", "12 new jobs for you"),
    "2": ("SEEK <jobs@s.seek.com.au>", "New job matching Data Scientist"),
    "3": ("LinkedIn <jobalerts-noreply@linkedin.com>", "Your job alert"),
    "4": ("Substack <news@substack.com>", "The Monday briefing"),
    "5": ("Acme <talent@acme-hiring.io>", "Engineering roles open"),
}
PAGES = {None: (["1", "2", "3"], "page2"), "page2": (["4", "5"], None)}


@pytest.fixture
def gmail(monkeypatch):
    service, box = _fake_service(MSGS, PAGES)
    monkeypatch.setattr(ingest, "_gmail_service", lambda: service)
    return box


@pytest.fixture(autouse=True)
def temp_settings(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "SETTINGS_FILE", tmp_path / "settings.json")


# --- From-header parsing -----------------------------------------------------

@pytest.mark.parametrize("raw,expected", [
    ("SEEK <noreply@s.Seek.COM.AU>", ("SEEK", "noreply@s.seek.com.au", "s.seek.com.au")),
    ('"Doe, Jane" <jane@example.com>', ("Doe, Jane", "jane@example.com", "example.com")),
    ("plain@example.org", ("", "plain@example.org", "example.org")),
    ("", ("", "", "")),
])
def test_parse_from(raw, expected):
    """parseaddr rather than a regex: display names legitimately contain commas
    and angle brackets, which is exactly where hand-rolled splitting breaks."""
    assert ingest._parse_from(raw) == expected


# --- scanning ----------------------------------------------------------------

def test_scan_paginates_and_groups_by_domain(gmail):
    found = ingest.scan_senders(max_messages=50)
    domains = {r["domain"]: r for r in found}

    assert len(gmail.list_calls) == 2, "second page was not requested"
    assert set(domains) == {"s.seek.com.au", "linkedin.com", "substack.com",
                            "acme-hiring.io"}
    # Boards rotate the local part but keep the domain, so the two SEEK addresses
    # must collapse into one row rather than competing for the same checkbox.
    assert domains["s.seek.com.au"]["count"] == 2
    assert found[0]["domain"] == "s.seek.com.au", "not sorted most-frequent-first"


def test_scan_requests_only_metadata_headers(gmail):
    ingest.scan_senders(max_messages=50)
    call = gmail.get_calls[0]
    assert call["format"] == "metadata"
    assert call["metadataHeaders"] == ["From", "Subject"]


def test_scan_collects_at_most_two_sample_subjects(gmail):
    found = {r["domain"]: r for r in ingest.scan_senders(max_messages=50)}
    assert found["s.seek.com.au"]["sample_subjects"] == [
        "12 new jobs for you", "New job matching Data Scientist"]


def test_scan_respects_max_messages(gmail):
    ingest.scan_senders(max_messages=2)
    assert len(gmail.get_calls) == 2


def test_scan_reports_progress(gmail):
    progress = {}
    ingest.scan_senders(max_messages=50, progress=progress)
    assert progress["scanned"] == 5


# --- classification ----------------------------------------------------------

@pytest.mark.parametrize("domain,address,subjects", [
    ("s.seek.com.au", "noreply@s.seek.com.au", []),
    ("linkedin.com", "jobalerts-noreply@linkedin.com", []),
    ("mail.example.com", "hi@mail.example.com", ["New jobs for you this week"]),
])
def test_heuristic_accepts_obvious_job_alerts(domain, address, subjects):
    assert ingest._looks_like_job_alert(domain, address, subjects) is True


@pytest.mark.parametrize("domain,address,subjects", [
    ("substack.com", "news@substack.com", ["The Monday briefing"]),
    ("bank.com.au", "statements@bank.com.au", ["Your statement is ready"]),
])
def test_heuristic_rejects_non_job_mail(domain, address, subjects):
    assert ingest._looks_like_job_alert(domain, address, subjects) is False


def test_undecided_senders_go_to_the_llm_in_one_batched_call(monkeypatch, gmail):
    """A busy inbox turns up a hundred domains; per-sender calls would make this
    slow and needlessly costly."""
    calls = []

    def fake_classify(rows):
        calls.append([r["domain"] for r in rows])
        return {"substack.com": False}

    monkeypatch.setattr(ingest, "_llm_classify", fake_classify)
    ingest.classify_senders(ingest.scan_senders(max_messages=50))

    assert len(calls) == 1, "classifier should be called once, not per sender"
    # Only what the heuristic could not place is sent.
    assert calls[0] == ["substack.com"]


def test_llm_suggestion_is_marked_distinctly(monkeypatch, gmail):
    monkeypatch.setattr(ingest, "_llm_classify", lambda rows: {"substack.com": True})
    out = {r["domain"]: r["confidence"]
           for r in ingest.classify_senders(ingest.scan_senders(max_messages=50))}
    assert out["s.seek.com.au"] == "heuristic"
    assert out["substack.com"] == "llm"


def test_classification_failure_degrades_to_unknown(monkeypatch, gmail):
    """No API key, or a transient error, must not sink the whole scan — those
    senders simply appear unticked, which is where they'd have been anyway."""
    def boom(rows):
        raise RuntimeError("no key")

    monkeypatch.setattr(ingest, "_llm_classify", boom)
    out = {r["domain"]: r["confidence"]
           for r in ingest.classify_senders(ingest.scan_senders(max_messages=50))}
    assert out["substack.com"] == "unknown"
    assert out["s.seek.com.au"] == "heuristic"      # heuristic results survive


def test_llm_classify_parses_verdicts(monkeypatch):
    monkeypatch.setattr(ingest, "_llm_classify", ingest._llm_classify)
    import stages.ingest as ing

    def fake_structured(system, user, schema, **kw):
        assert schema is SenderClassification
        return SenderClassification(verdicts=[
            SenderVerdict(domain="  Acme-Hiring.IO  ", is_job_alert=True)])

    monkeypatch.setattr("llm.call_structured", fake_structured)
    out = ing._llm_classify([{"domain": "acme-hiring.io", "address": "a@b.c",
                              "sample_subjects": []}])
    assert out == {"acme-hiring.io": True}, "domain should be normalised for lookup"


def test_discover_marks_already_configured_senders(monkeypatch, gmail):
    settings.set_sources([EmailSource(value="s.seek.com.au")])
    monkeypatch.setattr(ingest, "_llm_classify", lambda rows: {})
    found = {r["domain"]: r for r in ingest.discover_senders()}
    assert found["s.seek.com.au"]["already_added"] is True
    assert found["linkedin.com"]["already_added"] is False


# --- the query ingest actually runs ------------------------------------------

def test_fetch_builds_its_query_from_saved_senders(monkeypatch):
    settings.set_sources([EmailSource(value="seek.com.au"),
                          EmailSource(value="indeed.com", enabled=False)])
    msgs = {"1": ("SEEK <a@seek.com.au>", "jobs")}
    service, box = _fake_service(msgs, {None: (["1"], None)})
    monkeypatch.setattr(ingest, "_gmail_service", lambda: service)
    monkeypatch.setattr(ingest, "extract_text", lambda payload: "body text")

    ingest.fetch_job_emails(days=3)

    q = box.list_calls[0]["q"]
    assert q == "from:(seek.com.au) newer_than:3d"
    assert "indeed.com" not in q, "a disabled sender must not reach the query"


def test_fetch_paginates(monkeypatch):
    """This used to read only the first page and silently drop the rest — easy to
    miss while the sender list was a fixed handful."""
    msgs = {str(i): (f"S <a{i}@seek.com.au>", "jobs") for i in range(1, 6)}
    service, box = _fake_service(msgs, PAGES)
    monkeypatch.setattr(ingest, "_gmail_service", lambda: service)
    monkeypatch.setattr(ingest, "extract_text", lambda payload: "body text")

    results = ingest.fetch_job_emails(days=7)

    assert len(box.list_calls) == 2
    assert len(results) == 5, "messages beyond the first page were dropped"


def test_fetch_honours_an_explicit_query_override(monkeypatch):
    service, box = _fake_service({"1": ("S <a@b.com>", "x")}, {None: (["1"], None)})
    monkeypatch.setattr(ingest, "_gmail_service", lambda: service)
    monkeypatch.setattr(ingest, "extract_text", lambda payload: "body")
    ingest.fetch_job_emails(query="from:(custom.com)")
    assert box.list_calls[0]["q"] == "from:(custom.com)"


# --- fact bank extraction ----------------------------------------------------

def test_render_markdown_marks_uncertain_items():
    """The grounding check treats this file as truth, so an unmarked guess would
    silently license a fabricated resume claim."""
    doc = FactBankDoc(name="Jane Doe", contact="jane@example.com", sections=[
        FactSection(heading="Experience", items=[
            FactItem(text="Used scikit-learn on a churn dataset."),
            FactItem(text="Some cloud tooling.", uncertain=True)])])
    md = factbank.render_markdown(doc)

    assert "## Experience" in md
    assert "- Used scikit-learn on a churn dataset." in md
    assert "- Some cloud tooling.  *(verify)*" in md
    assert "REVIEW THIS BEFORE RELYING ON IT" in md


def test_extract_fact_bank_uses_the_quality_tier_at_zero_temperature(monkeypatch):
    """The cheap tier is markedly worse at resisting the urge to smooth vague
    phrasing into confident claims, and this file is the app's ground truth."""
    seen = {}

    def fake_structured(system, user, schema, tier="cheap", temperature=0.0, **kw):
        seen.update(tier=tier, temperature=temperature, user=user)
        return FactBankDoc(sections=[FactSection(
            heading="Skills", items=[FactItem(text="Python")])])

    monkeypatch.setattr(factbank, "call_structured", fake_structured)
    md = factbank.extract_fact_bank("Jane Doe\nSkills: Python")

    assert seen["tier"] == "quality"
    assert seen["temperature"] == 0.0
    assert "- Python" in md


def test_extract_rejects_empty_text():
    with pytest.raises(ValueError):
        factbank.extract_fact_bank("   ")


def test_extract_truncates_oversized_input(monkeypatch):
    """Guards against a 200-page PDF going to the quality tier in full."""
    seen = {}

    def fake_structured(system, user, schema, **kw):
        seen["len"] = len(user)
        return FactBankDoc(sections=[FactSection(
            heading="S", items=[FactItem(text="x")])])

    monkeypatch.setattr(factbank, "call_structured", fake_structured)
    factbank.extract_fact_bank("word " * 40_000)
    assert seen["len"] < factbank.MAX_CHARS + 200


def test_extract_docx_reads_table_cells():
    """Plenty of resume templates lay the whole document out in an invisible
    table; reading only paragraphs returns almost nothing for those."""
    import io

    import docx
    d = docx.Document()
    d.add_paragraph("Jane Doe")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Analyst Intern"
    table.rows[0].cells[1].text = "Used scikit-learn."
    buf = io.BytesIO()
    d.save(buf)

    text = factbank.extract_docx_text(buf.getvalue())
    assert "Jane Doe" in text and "Used scikit-learn." in text


def test_unsupported_extension_is_rejected():
    with pytest.raises(ValueError):
        factbank.extract_text(b"data", ".txt")
